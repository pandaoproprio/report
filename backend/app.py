from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from sqlalchemy.orm import Session
from datetime import datetime
from io import BytesIO
import os

from database import engine, SessionLocal, init_db
from models import Report, Base

from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

app = Flask(__name__)
CORS(app)

init_db()

@app.route('/')
def home():
    return jsonify({"message": "Sistema de Relatórios API", "version": "1.0.0"})

@app.route('/api/reports', methods=['GET'])
def get_reports():
    db = SessionLocal()
    try:
        category = request.args.get('category')
        status = request.args.get('status')

        query = db.query(Report)

        if category:
            query = query.filter(Report.category == category)
        if status:
            query = query.filter(Report.status == status)

        reports = query.order_by(Report.created_at.desc()).all()
        return jsonify([report.to_dict() for report in reports])
    finally:
        db.close()

@app.route('/api/reports', methods=['POST'])
def create_report():
    db = SessionLocal()
    try:
        data = request.json
        report = Report(
            title=data.get('title'),
            description=data.get('description'),
            category=data.get('category'),
            value=data.get('value', 0.0),
            status=data.get('status', 'pending')
        )
        db.add(report)
        db.commit()
        db.refresh(report)
        return jsonify(report.to_dict()), 201
    finally:
        db.close()

@app.route('/api/reports/<int:report_id>', methods=['GET'])
def get_report(report_id):
    db = SessionLocal()
    try:
        report = db.query(Report).filter(Report.id == report_id).first()
        if not report:
            return jsonify({"error": "Relatório não encontrado"}), 404
        return jsonify(report.to_dict())
    finally:
        db.close()

@app.route('/api/reports/<int:report_id>', methods=['PUT'])
def update_report(report_id):
    db = SessionLocal()
    try:
        report = db.query(Report).filter(Report.id == report_id).first()
        if not report:
            return jsonify({"error": "Relatório não encontrado"}), 404

        data = request.json
        if 'title' in data:
            report.title = data['title']
        if 'description' in data:
            report.description = data['description']
        if 'category' in data:
            report.category = data['category']
        if 'value' in data:
            report.value = data['value']
        if 'status' in data:
            report.status = data['status']

        report.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(report)
        return jsonify(report.to_dict())
    finally:
        db.close()

@app.route('/api/reports/<int:report_id>', methods=['DELETE'])
def delete_report(report_id):
    db = SessionLocal()
    try:
        report = db.query(Report).filter(Report.id == report_id).first()
        if not report:
            return jsonify({"error": "Relatório não encontrado"}), 404

        db.delete(report)
        db.commit()
        return jsonify({"message": "Relatório deletado com sucesso"})
    finally:
        db.close()

@app.route('/api/stats', methods=['GET'])
def get_stats():
    db = SessionLocal()
    try:
        total_reports = db.query(Report).count()

        by_status = {}
        statuses = db.query(Report.status).distinct().all()
        for (status,) in statuses:
            count = db.query(Report).filter(Report.status == status).count()
            by_status[status or 'unknown'] = count

        by_category = {}
        categories = db.query(Report.category).distinct().all()
        for (category,) in categories:
            count = db.query(Report).filter(Report.category == category).count()
            by_category[category or 'uncategorized'] = count

        total_value = db.query(Report).with_entities(Report.value).all()
        sum_value = sum([v[0] or 0 for v in total_value])

        return jsonify({
            "total_reports": total_reports,
            "by_status": by_status,
            "by_category": by_category,
            "total_value": round(sum_value, 2)
        })
    finally:
        db.close()

@app.route('/api/reports/export/pdf', methods=['GET'])
def export_pdf():
    db = SessionLocal()
    try:
        reports = db.query(Report).order_by(Report.created_at.desc()).all()

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#2c3e50'),
            spaceAfter=30,
            alignment=TA_CENTER
        )

        title = Paragraph("Relatório Geral", title_style)
        elements.append(title)
        elements.append(Spacer(1, 0.3*inch))

        data = [['ID', 'Título', 'Categoria', 'Status', 'Valor']]

        for report in reports:
            data.append([
                str(report.id),
                report.title[:30] if report.title else '',
                report.category or '',
                report.status or '',
                f'R$ {report.value:.2f}' if report.value else 'R$ 0.00'
            ])

        table = Table(data, colWidths=[0.7*inch, 2.5*inch, 1.5*inch, 1.2*inch, 1.2*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
        ]))

        elements.append(table)
        doc.build(elements)

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=f'relatorio_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf'
        )
    finally:
        db.close()

@app.route('/api/reports/export/excel', methods=['GET'])
def export_excel():
    db = SessionLocal()
    try:
        reports = db.query(Report).order_by(Report.created_at.desc()).all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Relatórios"

        headers = ['ID', 'Título', 'Descrição', 'Categoria', 'Status', 'Valor', 'Criado em']
        ws.append(headers)

        header_fill = PatternFill(start_color="3498db", end_color="3498db", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for report in reports:
            ws.append([
                report.id,
                report.title,
                report.description,
                report.category,
                report.status,
                report.value,
                report.created_at.strftime('%Y-%m-%d %H:%M:%S') if report.created_at else ''
            ])

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'relatorio_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )
    finally:
        db.close()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
